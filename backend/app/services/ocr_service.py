import io
import re
from pathlib import Path
from PIL import Image

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from docx import Document
except ImportError:
    Document = None

import pytesseract


def normalize_text(text: str) -> str:
    """
    Standardize text formatting:
    - Normalize line breaks and remove form feeds
    - Collapse 3+ newlines to 2 (preserving paragraph structure)
    - Collapse multiple spaces and tabs to a single space
    """
    if not text:
        return ""

    # Replace carriage returns and form feeds
    cleaned = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n")

    # Clean horizontal whitespace per line
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in cleaned.split("\n")]
    cleaned = "\n".join(lines)

    # Collapse excessive vertical whitespace to 2 newlines (paragraph boundary)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _extract_pdf_fitz(doc) -> str:
    """
    Extract text page-by-page using PyMuPDF.
    If a page has fewer than 50 words, fallback to high-DPI Tesseract OCR for that page.
    """
    page_texts: list[str] = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        raw_text = page.get_text("text") or ""
        words = raw_text.split()

        # If word count is low (< 50 words), page is likely a scanned sheet or image
        if len(words) < 50:
            try:
                pix = page.get_pixmap(dpi=300)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                ocr_text = pytesseract.image_to_string(img, lang="eng").strip()
                if ocr_text:
                    page_texts.append(ocr_text)
                else:
                    page_texts.append(raw_text)
            except Exception:
                page_texts.append(raw_text)
        else:
            page_texts.append(raw_text)

    return "\n\n".join(page_texts)


def extract_text(file_path: str, mime_type: str = "") -> str:
    """
    Extract text from a file on disk (PDF, DOCX, Images, Plain Text).
    """
    path = Path(file_path)
    if not path.is_file():
        return ""

    mime = (mime_type or "").lower()
    suffix = path.suffix.lower()

    is_pdf = "pdf" in mime or suffix == ".pdf"
    is_docx = "word" in mime or "docx" in mime or suffix == ".docx"
    is_image = "image" in mime or suffix in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]
    is_text = mime.startswith("text/") or "plain" in mime or suffix in [".txt", ".md", ".csv", ".json"]

    extracted_text = ""

    if is_pdf:
        if fitz is not None:
            try:
                with fitz.open(file_path) as doc:
                    extracted_text = _extract_pdf_fitz(doc)
            except Exception:
                extracted_text = ""
        else:
            # Fallback if PyMuPDF is not installed
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(file_path)
                extracted_text = "\n".join((p.extract_text() or "") for p in reader.pages)
            except Exception:
                extracted_text = ""

    elif is_docx:
        if Document is not None:
            try:
                document = Document(file_path)
                parts = [p.text for p in document.paragraphs if p.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            parts.append(row_text)
                extracted_text = "\n\n".join(parts)
            except Exception:
                extracted_text = ""

    elif is_image:
        try:
            with Image.open(file_path) as image:
                extracted_text = pytesseract.image_to_string(image, lang="eng")
        except Exception:
            extracted_text = ""

    elif is_text or not mime:
        try:
            extracted_text = path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            extracted_text = ""

    return normalize_text(extracted_text)


def extract_text_from_bytes(data: bytes, mime_type: str = "", filename: str = "") -> str:
    """
    Extract text directly from bytes (for direct streaming or R2 ingestion).
    """
    if not data:
        return ""

    mime = (mime_type or "").lower()
    suffix = Path(filename).suffix.lower() if filename else ""

    is_pdf = "pdf" in mime or suffix == ".pdf"
    is_docx = "word" in mime or "docx" in mime or suffix == ".docx"
    is_image = "image" in mime or suffix in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]
    is_text = mime.startswith("text/") or "plain" in mime or suffix in [".txt", ".md", ".csv", ".json"]

    extracted_text = ""

    if is_pdf:
        if fitz is not None:
            try:
                with fitz.open(stream=data, filetype="pdf") as doc:
                    extracted_text = _extract_pdf_fitz(doc)
            except Exception:
                extracted_text = ""
        else:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(io.BytesIO(data))
                extracted_text = "\n".join((p.extract_text() or "") for p in reader.pages)
            except Exception:
                extracted_text = ""

    elif is_docx:
        if Document is not None:
            try:
                document = Document(io.BytesIO(data))
                parts = [p.text for p in document.paragraphs if p.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            parts.append(row_text)
                extracted_text = "\n\n".join(parts)
            except Exception:
                extracted_text = ""

    elif is_image:
        try:
            with Image.open(io.BytesIO(data)) as image:
                extracted_text = pytesseract.image_to_string(image, lang="eng")
        except Exception:
            extracted_text = ""

    elif is_text or not mime:
        try:
            extracted_text = data.decode("utf-8", errors="ignore")
        except Exception:
            extracted_text = ""

    return normalize_text(extracted_text)
